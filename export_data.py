#!/usr/bin/env python3
"""Export Viral_Content_Tracker.xlsx -> JSON for the tracker site."""
import json, os, re
from openpyxl import load_workbook

wb = load_workbook("/home/hatch/workspace/your_files/Viral_Content_Tracker.xlsx")

def rows_of(sheet):
    ws = wb[sheet]
    out = []
    for row in ws.iter_rows(values_only=True):
        out.append([("" if c is None else str(c).strip()) for c in row])
    return out

accounts, news = [], []
cur_date = None
for r in rows_of("Viral Accounts"):
    if not any(r):
        continue
    joined = " ".join(r)
    m = re.search(r"VIRAL ACCOUNTS — (\d{4}-\d{2}-\d{2})", joined)
    if m:
        cur_date = m.group(1)
        continue
    if r[0] == "Date" or len(r) < 12:
        continue
    # columns: Date, Category, #, Handle, Link, Niche, Followers, Views, Video Length, Monetizable, Why Viral, Top Pick
    accounts.append({
        "date": r[0] or cur_date, "category": r[1], "n": r[2], "handle": r[3],
        "link": r[4], "niche": r[5], "followers": r[6], "views": r[7],
        "length": r[8], "monetizable": r[9], "why": r[10], "top": r[11] == "⭐",
    })

cur_date = None
for r in rows_of("Viral News"):
    if not any(r):
        continue
    joined = " ".join(r)
    m = re.search(r"VIRAL NEWS — (\d{4}-\d{2}-\d{2})", joined)
    if m:
        cur_date = m.group(1)
        continue
    if r[0] == "Date" or len(r) < 9:
        continue
    # Date, #, Headline, Summary, Type, Source, Viral Signal, Confidence, Top Pick
    news.append({
        "date": r[0] or cur_date, "n": r[1], "headline": r[2], "summary": r[3],
        "type": r[4], "source": r[5], "signal": r[6], "confidence": r[7], "top": r[8] == "⭐",
    })

trending = []
_tpath = "/home/hatch/workspace/tracker-site/trending.json"
if os.path.exists(_tpath):
    trending = json.load(open(_tpath))

json.dump({"accounts": accounts, "news": news, "trending": trending},
          open("/home/hatch/workspace/tracker-site/data.json", "w"), ensure_ascii=False, indent=1)
print(f"accounts={len(accounts)} news={len(news)} trending={len(trending)}")
dates = sorted({a['date'] for a in accounts}, reverse=True)
print("account dates:", dates)
niches = sorted({a['niche'] for a in accounts})
print("niches:", niches)
